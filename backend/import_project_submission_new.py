"""Import the selected A-Inno 2569 workbook into ProjectSubmissionNew.

The command is non-mutating unless --apply is passed.  Every applied import uses
data_source='a-inno' and emits a JSON report that can be used for rollback.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from openpyxl import load_workbook


SOURCE_FIELD_NAMES = [
    "creativityID", "year", "creativityName", "submission_type", "topic1",
    "topic2", "topic2_so", "topic3", "topic4", "topic4_detail", "topic5",
    "topic6", "topic7", "topic8", "topic9", "topic9_1", "topic9_2",
    "topic9_3", "topic9_4", "creativityStatusTxt", "SendBy_EmpID",
    "SendBy_fullname", "SendBy_division_name", "SendBy_contact", "teamdetail",
] + [field for seq in range(1, 11) for field in (f"EmpCode{seq}", f"FullName{seq}", f"division_name{seq}")]

CHALLENGES = [
    "สนับสนุนการให้บริการให้สินเชื่อในภาค/นอกภาคการเกษตร", "สนับสนุนการให้บริการทางการเงิน",
    "กระบวนการติดตามการชำระหนี้", "เพิ่มโอกาสธุรกิจประกันภัย", "การลดการใช้พลังงาน",
    "การจัดการข้อมูลลูกค้า", "การเพิ่มรายได้ FBI", "การบริหารจัดการโครงการนโยบายรัฐ",
    "การเข้าถึงบริการทางการเงินของผู้สูงอายุ", "เพิ่มมูลค่าจากบริการ mobile digital",
    "ช่องทางการเข้าถึงตลาดสินค้าเกษตร", "การยกระดับชุมชนของธนาคาร", "การสร้างรายได้ใหม่ให้เกษตรกร",
]
CHALLENGE_CATEGORIES = ["เงินฝาก", "FBI", "การเติบโตของสินเชื่อ (Growth)", "การบริหารจัดการหนี้", "การพัฒนาลูกค้า", "การพัฒนาปรับปรุงกระบวนการ", "การลดค่าใช้จ่ายของส่วนงาน", "การดำเนินธุรกิจอย่างยั่งยืน (ESG)"]
INNOVATION_TYPES = ["นวัตกรรมผลิตภัณฑ์", "นวัตกรรมบริการ", "นวัตกรรมกระบวนการ", "นวัตกรรมรูปแบบการดำเนินธุรกิจหรือภารกิจใหม่ขององค์กร", "นวัตกรรมเกษตรที่ส่งเสริมการยกระดับอาชีพเกษตร"]
DIGITAL_TYPES = ["เป็นนวัตกรรมที่ใช้เทคโนโลยีดิจิทัล", "เป็นนวัตกรรมที่ไม่ใช้เทคโนโลยีดิจิทัล"]
NOVELTY_LEVELS = [
    "การปรับปรุงผลิตภัณฑ์ บริการ กระบวนการ หรือองค์ความรู้เดิม",
    "การพัฒนา/ต่อยอดผลิตภัณฑ์ บริการ กระบวนการ หรือองค์ความรู้เดิม",
    "การคิดสร้างสรรค์ผลิตภัณฑ์ บริการ กระบวนการ หรือองค์ความรู้ใหม่ระดับองค์กร",
    "การคิดสร้างสรรค์ผลิตภัณฑ์ บริการ กระบวนการ หรือองค์ความรู้ใหม่ระดับอุตสาหกรรม/ธุรกิจ",
    "การคิดสร้างสรรค์ผลิตภัณฑ์ บริการ กระบวนการ หรือองค์ความรู้ใหม่ระดับประเทศ",
]
IDEA_SOURCE_KEYS = {"ชุมชนนักปฏิบัติ": "CoPs", "คลังความรู้": "LR", "ผลงานวิจัย": "Research", "ประสบการณ์": "Experience", "ศึกษาดูงาน": "StudyVisit", "แลกเปลี่ยนเรียนรู้": "KnowledgeExchange", "ฐานข้อมูลนวัตกรรม": "InnovationDatabase", "ศึกษาความต้องการของตลาด": "MarketStudy", "VOS": "VOS"}


def clean(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def comparable(value: str) -> str:
    return re.sub(r"\s+", "", clean(value))


def normalize_idea_name(value: Any) -> str:
    """Ignore spaces and punctuation while preserving Unicode letters and digits."""
    normalized = unicodedata.normalize("NFKC", clean(value)).casefold()
    return "".join(character for character in normalized if character.isalnum())


def number_and_text(value: str, choices: list[str], field: str, excel_row: int) -> tuple[int | None, str | None]:
    value = clean(value)
    if not value:
        return None, None
    for number, label in enumerate(choices, 1):
        if comparable(value) == comparable(label):
            return number, label
    raise ValueError(f"Excel row {excel_row}: unknown {field}: {value!r}")


def as_html(value: str) -> str | None:
    value = clean(value)
    return html.escape(value).replace("\n", "<br>") if value else None


def checked(value: str, marker: str) -> bool:
    return marker in clean(value)


def recover_shifted_members(team_detail: str, org_name: str) -> list[dict[str, Any]]:
    """Recover the one legacy export layout that stored members in SendBy_contact."""
    pattern = re.compile(r"([^,#]+)#(\d{1,20})#เบอร์ติดต่อ-([^,]+)")
    members = []
    for name, code, mobile in pattern.findall(team_detail):
        members.append({
            "MemberSeq": len(members) + 1,
            "EmpCode": code,
            "FullNameTh": name.strip(),
            "PositionName": None,
            "OrgName": org_name or None,
            "MobileNo": mobile.strip() or None,
            "IsTeamLeader": len(members) == 0,
            "IsMainContact": len(members) == 0,
        })
    return members


def map_source_row(row: dict[str, Any], excel_row: int) -> SimpleNamespace:
    event_year = clean(row.get("year"))
    idea_name = clean(row.get("creativityName"))
    submission_type = clean(row.get("submission_type") or row.get("รูปแบบ"))
    if not event_year.isdigit() or not idea_name:
        raise ValueError(f"Excel row {excel_row}: year and creativityName are required")
    submission_types = {"รายบุคคล": ("INDIVIDUAL", "ประเภทบุคคล"), "รายทีม": ("TEAM", "ประเภททีม")}
    if submission_type not in submission_types:
        raise ValueError(f"Excel row {excel_row}: unknown submission type: {submission_type!r}")
    challenge_no, challenge_text = number_and_text(row.get("topic1", ""), CHALLENGES, "topic1", excel_row)
    category_no, category_text = number_and_text(row.get("topic2", ""), CHALLENGE_CATEGORIES, "topic2", excel_row)
    innovation_no, innovation_text = number_and_text(row.get("topic5", ""), INNOVATION_TYPES, "topic5", excel_row)
    digital_no, digital_text = number_and_text(row.get("topic7", ""), DIGITAL_TYPES, "topic7", excel_row)
    novelty_no, novelty_text = number_and_text(row.get("topic8", ""), NOVELTY_LEVELS, "topic8", excel_row)
    customer_text = clean(row.get("topic4"))
    customer_no = {"ลูกค้าเดิม": 1, "ลูกค้าใหม่": 2}.get(customer_text)
    project = {
        "EventYear": int(event_year), "SubmissionTypeCode": submission_types[submission_type][0],
        "SubmissionTypeNameTh": submission_types[submission_type][1], "TeamName": None,
        "CreativeIdeaName": idea_name, "ChallengeNo": challenge_no, "ChallengeText": challenge_text,
        "ChallengeCategoryNo": category_no, "ChallengeCategoryText": category_text,
        "TargetCustomerTypeNo": customer_no, "TargetCustomerTypeText": customer_text or None,
        "TargetCustomerProblemHtml": as_html(row.get("topic4_detail", "")), "InnovationTypeNo": innovation_no,
        "InnovationTypeText": innovation_text, "IdeaConceptHtml": as_html(row.get("topic6", "")),
        "DigitalInnovationNo": digital_no, "DigitalInnovationText": digital_text,
        "NoveltyLevelNo": novelty_no, "NoveltyLevelText": novelty_text,
        "InnovationValueFinancial": checked(row.get("topic9", ""), "ด้านการเงิน"),
        "FinancialValueRevenue": checked(row.get("topic9_1", ""), "สร้างรายได้"),
        "FinancialValueCostSaving": checked(row.get("topic9_1", ""), "ลดค่าใช้จ่าย"),
        "FinancialValueDetailHtml": as_html(row.get("topic9_2", "")),
        "InnovationValueNonFinancial": checked(row.get("topic9", ""), "ด้านไม่ใช่การเงิน"),
        "NonFinancialValueCustomerSatisfaction": checked(row.get("topic9_3", ""), "ความพึงพอใจ"),
        "NonFinancialValueWorkEfficiency": checked(row.get("topic9_3", ""), "ลดขั้นตอนการทำงาน"),
        "NonFinancialValueCustomerQuality": checked(row.get("topic9_3", ""), "ด้านคุณภาพชีวิตลูกค้า"),
        "NonFinancialValueEnvironment": checked(row.get("topic9_3", ""), "ด้านสิ่งแวดล้อม"),
        "NonFinancialValueDetailHtml": as_html(row.get("topic9_4", "")), "StatusCode": "SUBMITTED",
        "CreatedByEmpCode": clean(row.get("SendBy_EmpID")) or None, "data_source": "a-inno",
    }
    selected_so = set(re.findall(r"SO([1-6])", clean(row.get("topic2_so"))))
    project.update({f"StrategicObjectiveSO{seq}": str(seq) in selected_so for seq in range(1, 7)})
    source_value = clean(row.get("topic3"))
    mapped_sources = []
    for marker, key in IDEA_SOURCE_KEYS.items():
        is_selected = marker in source_value
        project[f"IdeaSource{key}"] = is_selected
        project[f"IdeaSource{key}Detail"] = source_value if is_selected else None
        if is_selected:
            mapped_sources.append(key)
    project["IdeaSourceOther"] = not mapped_sources or "อื่น" in source_value
    project["IdeaSourceOtherDetail"] = source_value if project["IdeaSourceOther"] else None
    members = []
    for seq in range(1, 11):
        code, name, org = (clean(row.get(f"{prefix}{seq}")) for prefix in ("EmpCode", "FullName", "division_name"))
        if not any((code, name, org)):
            continue
        if not code or not name:
            raise ValueError(f"Excel row {excel_row}: member {seq} requires EmpCode and FullName")
        mobile_no = clean(row.get("SendBy_contact")) or None if seq == 1 else None
        members.append({"MemberSeq": len(members) + 1, "EmpCode": code, "FullNameTh": name, "PositionName": None, "OrgName": org or None, "MobileNo": mobile_no, "IsTeamLeader": len(members) == 0, "IsMainContact": len(members) == 0})
    if not members:
        raise ValueError(f"Excel row {excel_row}: at least one member is required")
    if not re.fullmatch(r"\d{1,20}", members[0]["EmpCode"]) and "#" in clean(row.get("SendBy_contact")):
        recovered = recover_shifted_members(clean(row.get("SendBy_contact")), clean(row.get("SendBy_division_name")))
        if recovered:
            members = recovered
    for member in members:
        for field, limit in {"EmpCode": 20, "FullNameTh": 200, "OrgName": 300, "MobileNo": 50}.items():
            value = member.get(field)
            if value is not None and len(value) > limit:
                raise ValueError(f"Excel row {excel_row}: {field} exceeds {limit} characters")
    return SimpleNamespace(project=project, members=members)


def read_source_rows(path: Path) -> list[tuple[int, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if len(headers) < len(SOURCE_FIELD_NAMES) or headers[:3] != ["creativityID", "year", "creativityName"]:
        raise ValueError("Unexpected workbook header; expected the selected A-Inno export format")
    return [(excel_row, dict(zip(SOURCE_FIELD_NAMES, values[:len(SOURCE_FIELD_NAMES)]))) for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2) if any(clean(value) for value in values[:len(SOURCE_FIELD_NAMES)])]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_import_plan(mapped_rows, existing_projects):
    existing_by_name = {}
    for project_id, idea_name, data_source in existing_projects:
        key = normalize_idea_name(idea_name)
        if key:
            existing_by_name.setdefault(key, []).append(
                {"ProjectId": project_id, "CreativeIdeaName": idea_name, "data_source": data_source}
            )

    first_source_row_by_name = {}
    accepted = []
    report_rows = []
    duplicate_source_count = 0
    duplicate_existing_count = 0
    empty_normalized_count = 0
    for excel_row, source, mapped in mapped_rows:
        idea_name = mapped.project["CreativeIdeaName"]
        key = normalize_idea_name(idea_name)
        report_row = {
            "excel_row": excel_row,
            "creativityID": clean(source.get("creativityID")),
            "CreativeIdeaName": idea_name,
            "member_count": len(mapped.members),
        }
        if not key:
            report_row.update({"status": "skipped", "skip_reason": "creative_idea_name_empty_after_normalization"})
            empty_normalized_count += 1
        elif key in first_source_row_by_name:
            report_row.update({
                "status": "skipped",
                "skip_reason": "duplicate_in_source_file",
                "duplicate_of_excel_row": first_source_row_by_name[key],
            })
            duplicate_source_count += 1
        elif key in existing_by_name:
            report_row.update({
                "status": "skipped",
                "skip_reason": "duplicate_existing_creative_idea_name",
                "matched_projects": existing_by_name[key],
            })
            first_source_row_by_name[key] = excel_row
            duplicate_existing_count += 1
        else:
            report_row["status"] = "planned"
            first_source_row_by_name[key] = excel_row
            accepted.append((excel_row, source, mapped, report_row))
        report_rows.append(report_row)

    summary = {
        "source_rows": len(mapped_rows),
        "planned_projects": len(accepted),
        "planned_members": sum(len(mapped.members) for _, _, mapped, _ in accepted),
        "duplicate_source_rows_skipped": duplicate_source_count,
        "duplicate_existing_rows_skipped": duplicate_existing_count,
        "empty_normalized_names_skipped": empty_normalized_count,
    }
    return accepted, report_rows, summary


def write_report(path: Path, report: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--report", type=Path, default=Path("import-reports/a_inno_2569.json"))
    args = parser.parse_args()
    mapped_rows = [
        (excel_row, source, map_source_row(source, excel_row))
        for excel_row, source in read_source_rows(args.input)
    ]
    from sqlalchemy import text
    from app.db.database import SessionLocal
    from app.db import models

    session = SessionLocal()
    try:
        existing_projects = (
            session.query(
                models.ProjectSubmissionNew.ProjectId,
                models.ProjectSubmissionNew.CreativeIdeaName,
                models.ProjectSubmissionNew.data_source,
            )
            .filter(models.ProjectSubmissionNew.data_source != "a-inno")
            .all()
        )
        a_inno_projects_to_delete = session.query(models.ProjectSubmissionNew.ProjectId).filter(
            models.ProjectSubmissionNew.data_source == "a-inno"
        ).all()
        old_project_ids = [row[0] for row in a_inno_projects_to_delete]
        old_project_count = len(old_project_ids)
        old_member_count = (
            session.query(models.ProjectSubmissionNewMember.ProjectMemberId)
            .filter(models.ProjectSubmissionNewMember.ProjectId.in_(old_project_ids))
            .count()
            if old_project_ids else 0
        )
        accepted, report_rows, summary = build_import_plan(mapped_rows, existing_projects)
        if args.apply and not accepted:
            raise RuntimeError("No unique A-Inno projects remain after duplicate checks; refusing to delete existing a-inno data")
        run_id = datetime.now().strftime("%Y%m%d%H%M%S%f")
        parent_backup = f"ProjectSubmissionNewAInnoBackup_{run_id}"
        member_backup = f"ProjectSubmissionNewMemberAInnoBackup_{run_id}"
        report = {
            "input": str(args.input),
            "sha256": sha256(args.input),
            "mode": "apply" if args.apply else "dry-run",
            "status": "pending" if args.apply else "planned",
            "data_source": "a-inno",
            "duplicate_match_rule": "NFKC + casefold; remove every non-letter and non-digit Unicode character",
            "summary": summary,
            "a_inno_existing_projects_to_delete": old_project_count,
            "a_inno_existing_members_to_delete": old_member_count,
            "backup_tables": {"projects": f"dbo.{parent_backup}", "members": f"dbo.{member_backup}"} if args.apply else None,
            "rows": report_rows,
        }

        if args.apply:
            write_report(args.report, report)
            try:
                session.execute(text(
                    f"SELECT * INTO dbo.[{parent_backup}] "
                    "FROM dbo.ProjectSubmissionNew WHERE data_source = :data_source"
                ), {"data_source": "a-inno"})
                session.execute(text(
                    f"SELECT m.* INTO dbo.[{member_backup}] "
                    "FROM dbo.ProjectSubmissionNewMember AS m "
                    "INNER JOIN dbo.ProjectSubmissionNew AS project "
                    "ON project.ProjectId = m.ProjectId "
                    "WHERE project.data_source = :data_source"
                ), {"data_source": "a-inno"})
                session.execute(text(
                    "DELETE m FROM dbo.ProjectSubmissionNewMember AS m "
                    "INNER JOIN dbo.ProjectSubmissionNew AS project "
                    "ON project.ProjectId = m.ProjectId "
                    "WHERE project.data_source = :data_source"
                ), {"data_source": "a-inno"})
                session.execute(text(
                    "DELETE FROM dbo.ProjectSubmissionNew WHERE data_source = :data_source"
                ), {"data_source": "a-inno"})

                now = datetime.now()
                for _, _, item, report_row in accepted:
                    project = models.ProjectSubmissionNew(**item.project, CreatedAt=now, SubmittedAt=now)
                    session.add(project)
                    session.flush()
                    for member in item.members:
                        session.add(models.ProjectSubmissionNewMember(
                            ProjectId=project.ProjectId,
                            CreatedAt=now,
                            **member,
                        ))
                    report_row["ProjectId"] = project.ProjectId
                    report_row["status"] = "inserted"
                session.commit()
            except Exception as exc:
                session.rollback()
                report["status"] = "rolled_back"
                report["error"] = str(exc)
                write_report(args.report, report)
                raise
            report["status"] = "complete"
        else:
            report["backup_tables"] = None

        write_report(args.report, report)
        output = {
            "mode": report["mode"],
            "status": report["status"],
            **summary,
            "a_inno_existing_projects_to_delete": old_project_count,
            "a_inno_existing_members_to_delete": old_member_count,
            "report": str(args.report),
        }
        print(json.dumps(output, ensure_ascii=False))
    finally:
        session.close()


if __name__ == "__main__":
    main()
